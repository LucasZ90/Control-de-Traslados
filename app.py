from datetime import datetime, date
import io
import json
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import streamlit as st
from supabase import create_client

# Configuración de la página en español
st.set_page_config(
    page_title="Gestión de Traslados - Diálisis", page_icon="🚗", layout="centered"
)

# --- CONEXIÓN A SUPABASE ---
try:
  supabase_url = st.secrets["SUPABASE_URL"].strip().rstrip("/")
  supabase_key = st.secrets["SUPABASE_KEY"].strip()
  supabase = create_client(supabase_url, supabase_key)
except Exception as e:
  st.error(
      "⚠️ Error de configuración: Faltan las credenciales de Supabase en los"
      f" Secrets de Streamlit. Detalle: {e}"
  )
  st.stop()

# --- ESTILOS VISUALES MINIMALISTAS Y TIPOGRAFÍA DE ALTA CALIDAD ---
st.markdown("""
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@400;500;600;700&display=swap');

    html, body, [class*="css"], .stApp {
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        background-color: #0D1117;
        color: #E6EDF3;
    }
    
    h1 {
        font-size: 1.6rem !important;
        font-weight: 700 !important;
        color: #F0F6FC !important;
        margin-bottom: 0.2rem !important;
        letter-spacing: -0.02em;
    }

    h2, h3 {
        font-size: 1.15rem !important;
        font-weight: 600 !important;
        color: #C9D1D9 !important;
        margin-top: 1rem !important;
        margin-bottom: 0.5rem !important;
    }

    /* Tarjetas de pacientes estilo dashboard minimalista */
    [data-testid="stVerticalBlockBorderWrapper"] {
        background-color: #161B22 !important;
        border: 1px solid #30363D !important;
        border-radius: 12px !important;
        padding: 0.8rem 1rem !important;
        margin-bottom: 0.6rem !important;
        box-shadow: 0 2px 8px rgba(0,0,0,0.2);
    }

    /* Botones sutiles y modernos */
    div.stButton > button {
        border-radius: 8px !important;
        font-weight: 600 !important;
        padding: 0.35rem 0.2rem !important;
        font-size: 0.82rem !important;
        border: 1px solid #30363D !important;
        background-color: #21262D !important;
        color: #C9D1D9 !important;
        transition: all 0.2s ease-in-out;
    }

    div.stButton > button:hover {
        border-color: #58A6FF !important;
        color: #58A6FF !important;
        background-color: #161B22 !important;
    }

    /* Links tipo píldoras/chips */
    .chip-link {
        display: inline-block;
        background-color: #21262D;
        border: 1px solid #30363D;
        border-radius: 20px;
        padding: 2px 10px;
        font-size: 0.78rem;
        color: #58A6FF !important;
        text-decoration: none;
        margin-right: 6px;
        margin-bottom: 4px;
        font-weight: 500;
    }

    /* Expander limpio */
    .streamlit-expanderHeader {
        font-weight: 600 !important;
        font-size: 0.95rem !important;
        color: #F0F6FC !important;
    }
    
    /* Disminuir margen de columnas móviles */
    [data-testid="column"] {
        width: fit-content !important;
        flex: 1 !important;
        min-width: 0 !important;
    }
    </style>
""", unsafe_allow_html=True)

TARIFA_COMPLETO = 5000
TARIFA_MEDIO = 2500
TURNOS_DISPONIBLES = ["05:40", "08:40"]

DEFAULT_MWF = [
    {"nombre": "Duarte Catalina", "turno": "05:40", "maps_link": "", "contactos": []},
    {"nombre": "Ledesma Tatiana", "turno": "05:40", "maps_link": "", "contactos": []},
    {"nombre": "Mendoza Ramón", "turno": "05:40", "maps_link": "", "contactos": []},
    {"nombre": "Meza Adriana", "turno": "08:40", "maps_link": "", "contactos": []},
    {"nombre": "Paiva Oscar", "turno": "08:40", "maps_link": "", "contactos": []},
]

DEFAULT_TTS = [
    {"nombre": "Cantero Miguel", "turno": "05:40", "maps_link": "", "contactos": []},
    {"nombre": "Martinez José", "turno": "05:40", "maps_link": "", "contactos": []},
    {"nombre": "Píriz Maximiliano", "turno": "05:40", "maps_link": "", "contactos": []},
    {"nombre": "Ramírez Orlando", "turno": "08:40", "maps_link": "", "contactos": []},
    {"nombre": "Villalba Nelida", "turno": "08:40", "maps_link": "", "contactos": []},
    {"nombre": "Cardozo José", "turno": "08:40", "maps_link": "", "contactos": []},
]


@st.cache_data(ttl=600)
def cargar_pacientes_db():
  try:
    response = supabase.table("pacientes").select("*").execute()
    data = response.data
    if not data:
      for p in DEFAULT_MWF:
        supabase.table("pacientes").insert({
            "grupo": "MWF",
            "nombre": p["nombre"],
            "turno": p["turno"],
            "maps_link": p["maps_link"],
            "contactos": p["contactos"],
        }).execute()
      for p in DEFAULT_TTS:
        supabase.table("pacientes").insert({
            "grupo": "TTS",
            "nombre": p["nombre"],
            "turno": p["turno"],
            "maps_link": p["maps_link"],
            "contactos": p["contactos"],
        }).execute()
      response = supabase.table("pacientes").select("*").execute()
      data = response.data
    return data
  except Exception as e:
    st.error(f"Error al conectar con la tabla pacientes: {e}")
    return DEFAULT_MWF + DEFAULT_TTS


def guardar_paciente_db(p_id, grupo, nombre, turno, maps_link, contactos):
  try:
    if p_id:
      supabase.table("pacientes").update({
          "grupo": grupo,
          "nombre": nombre,
          "turno": turno,
          "maps_link": maps_link,
          "contactos": contactos,
      }).eq("id", p_id).execute()
    else:
      supabase.table("pacientes").insert({
          "grupo": grupo,
          "nombre": nombre,
          "turno": turno,
          "maps_link": maps_link,
          "contactos": contactos,
      }).execute()
    st.cache_data.clear()
    return True, None
  except Exception as e:
    return False, str(e)


def eliminar_paciente_db(p_id):
  try:
    if p_id:
      supabase.table("pacientes").delete().eq("id", p_id).execute()
    st.cache_data.clear()
    return True, None
  except Exception as e:
    return False, str(e)


@st.cache_data(ttl=300)
def cargar_registros_db():
  try:
    response = supabase.table("registros").select("*").execute()
    data = response.data
    datos = {}
    if data:
      for row in data:
        f_str = row["fecha"]
        try:
          if datetime.strptime(f_str, "%d/%m/%Y").date() >= date(2026, 10, 14):
            if f_str not in datos:
              datos[f_str] = {}
            datos[f_str][row["nombre"]] = {
                "ida": row["ida"],
                "vuelta": row["vuelta"],
                "listo": row["listo"],
                "estado": row["estado"],
                "valor": row["valor"],
                "turno_dia": row.get("turno_dia", "05:40"),
            }
        except ValueError:
          pass
    return datos
  except Exception:
    return {}


def guardar_registro_db(fecha_str, nombre, info):
  try:
    existing = (
        supabase.table("registros")
        .select("id")
        .eq("fecha", fecha_str)
        .eq("nombre", nombre)
        .execute()
    )
    payload = {
        "fecha": fecha_str,
        "nombre": nombre,
        "ida": info["ida"],
        "vuelta": info["vuelta"],
        "listo": info["listo"],
        "estado": info["estado"],
        "valor": info["valor"],
        "turno_dia": info.get("turno_dia", "05:40"),
    }
    if existing.data:
      supabase.table("registros").update(payload).eq(
          "id", existing.data[0]["id"]
      ).execute()
    else:
      supabase.table("registros").insert(payload).execute()
    st.cache_data.clear()
  except Exception as e:
    st.error(f"Error al guardar registro: {e}")


def eliminar_registro_db(fecha_str, nombre):
  try:
    supabase.table("registros").delete().eq("fecha", fecha_str).eq(
        "nombre", nombre
    ).execute()
    st.cache_data.clear()
  except Exception as e:
    st.error(f"Error al eliminar registro: {e}")


def actualizar_valores_paciente(info):
  reg_ida = info.get("ida", False)
  reg_vuelta = info.get("vuelta", False)
  if reg_ida and reg_vuelta:
    info["estado"] = "Completo (Ida y Vuelta)"
    info["valor"] = TARIFA_COMPLETO
  elif reg_ida or reg_vuelta:
    info["estado"] = "Medio (Un solo tramo)"
    info["valor"] = TARIFA_MEDIO
  else:
    info["estado"] = "Ausente (No viajó)"
    info["valor"] = 0


def generar_excel_bytes(datos, es_diario=False):
  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = "Resumen Diario" if es_diario else "Resumen Mensual"

  encabezados = ["Fecha", "Paciente", "Estado", "Valor ($)"]
  ws.append(encabezados)

  header_fill = PatternFill(
      start_color="1F4E78", end_color="1F4E78", fill_type="solid"
  )
  header_font = Font(color="FFFFFF", bold=True, size=11)
  thin_border = Border(
      left=Side(style="thin", color="D3D3D3"),
      right=Side(style="thin", color="D3D3D3"),
      top=Side(style="thin", color="D3D3D3"),
      bottom=Side(style="thin", color="D3D3D3"),
  )

  for col_num, header in enumerate(encabezados, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border

  fila_actual = 2
  for fecha, pacientes_dia in sorted(datos.items()):
    try:
      if datetime.strptime(fecha, "%d/%m/%Y").date() < date(2026, 10, 14):
        continue
    except ValueError:
      continue

    for nombre, info in pacientes_dia.items():
      ws.cell(row=fila_actual, column=1, value=fecha).border = thin_border
      ws.cell(row=fila_actual, column=2, value=nombre).border = thin_border
      ws.cell(row=fila_actual, column=3, value=info["estado"]).border = (
          thin_border
      )

      valor_num = (
          int(info["valor"])
          if isinstance(info["valor"], (int, float, str))
          and str(info["valor"]).isdigit()
          else 0
      )
      celda_val = ws.cell(row=fila_actual, column=4, value=valor_num)
      celda_val.border = thin_border
      celda_val.number_format = '"$"#,##0'

      fila_actual += 1

  ws.cell(row=fila_actual, column=3, value="TOTAL GENERAL:").font = Font(
      bold=True
  )
  celda_total = ws.cell(
      row=fila_actual, column=4, value=f"=SUM(D2:D{fila_actual-1})"
  )
  celda_total.font = Font(bold=True, color="1F4E78")
  celda_total.border = thin_border
  celda_total.number_format = '"$"#,##0'

  column_widths = {"A": 15, "B": 25, "C": 30, "D": 15}
  for col, width in column_widths.items():
    ws.column_dimensions[col].width = width

  output = io.BytesIO()
  wb.save(output)
  output.seek(0)
  return output.getvalue()


# Título principal limpio
st.title("Gestión de Traslados")

# --- NAVEGACIÓN SUPERIOR REDUCIDA ---
if "menu_actual" not in st.session_state:
  st.session_state["menu_actual"] = "Registrar / Actualizar Viajes"

col_nav1, col_nav2 = st.columns(2)
with col_nav1:
  if st.button("📋 Registrar", use_container_width=True):
    st.session_state["menu_actual"] = "Registrar / Actualizar Viajes"
with col_nav2:
  if st.button("📊 Resumen", use_container_width=True):
    st.session_state["menu_actual"] = "Resumen Mensual"

menu = st.session_state["menu_actual"]
st.markdown("---")

datos = cargar_registros_db()
todos_pacientes = cargar_pacientes_db()

if menu == "Registrar / Actualizar Viajes":
  st.subheader("📋 Registro Diario")

  # Botón superior para agregar nuevo paciente directamente
  with st.expander("➕ Agregar Nuevo Paciente al Sistema"):
    nuevo_nombre = st.text_input(
        "Nombre y Apellido del Paciente", key="nuevo_nombre_paciente"
    )
    nuevo_grupo = st.selectbox(
        "Grupo por Defecto",
        [
            "Grupo 1 (Lunes, Miércoles, Viernes)",
            "Grupo 2 (Martes, Jueves, Sábados)",
        ],
        key="nuevo_grupo_paciente",
    )
    nuevo_turno = st.selectbox(
        "Horario Habitual", TURNOS_DISPONIBLES, key="nuevo_turno_paciente"
    )
    nuevo_maps = st.text_input(
        "Enlace de Google Maps (Opcional)", key="nuevo_maps_paciente"
    )
    nuevo_tel = st.text_input(
        "Número de Teléfono Principal (Opcional)", key="nuevo_tel_paciente"
    )
    nueva_rel_tel = st.text_input(
        "Relación de este teléfono (Ej: Paciente, Hijo/a)",
        value="Paciente",
        key="nuevo_rel_paciente",
    )

    if st.button(
        "💾 Guardar Nuevo Paciente",
        type="primary",
        use_container_width=True,
        key="btn_guardar_nuevo_paciente",
    ):
      if nuevo_nombre.strip() == "":
        st.error("Por favor, escribe el nombre del paciente.")
      else:
        lista_contactos = []
        if nuevo_tel.strip():
          lista_contactos.append({
              "numero": nuevo_tel.strip(),
              "relacion": (
                  nueva_rel_tel.strip()
                  if nueva_rel_tel.strip()
                  else "Paciente"
              ),
          })

        nuevo_grupo_val = "MWF" if "Grupo 1" in nuevo_grupo else "TTS"
        ok, err = guardar_paciente_db(
            None,
            nuevo_grupo_val,
            nuevo_nombre.strip(),
            nuevo_turno,
            nuevo_maps.strip(),
            lista_contactos,
        )
        if ok:
          st.success(f"¡Paciente '{nuevo_nombre}' agregado con éxito!")
          st.rerun()
        else:
          st.error(f"Error al guardar nuevo paciente: {err}")

  fecha_seleccionada = st.date_input("Selecciona la fecha", datetime.now())
  fecha_str = fecha_seleccionada.strftime("%d/%m/%Y")

  if fecha_seleccionada < date(2026, 10, 14):
    st.error(
        "⚠️ Los registros de traslados comienzan a partir del 14 de octubre de"
        " 2026."
    )
  else:
    dia_semana = fecha_seleccionada.weekday()

    if dia_semana == 6:
      st.warning("⚠️ Los domingos no hay servicio programado.")
    else:
      nombre_dia = [
          "Lunes",
          "Martes",
          "Miércoles",
          "Jueves",
          "Viernes",
          "Sábado",
          "Domingo",
      ][dia_semana]
      grupo_sugerido = "MWF" if dia_semana in [0, 2, 4] else "TTS"
      st.info(
          f"📅 **{fecha_str}** ({nombre_dia}) — *Sugerencia: Grupo"
          f" {'1 (MWF)' if grupo_sugerido == 'MWF' else '2 (TTS)'}*"
      )

      if fecha_str not in datos:
        datos[fecha_str] = {}

      # --- CONFIGURACIÓN ESTRICTA DE PACIENTES PARA ESTE DÍA ---
      with st.expander(
          "⚙️ Seleccionar pacientes y turnos de hoy", expanded=False
      ):
        st.write(
            "Marca únicamente los pacientes que viajan hoy y ajusta su turno:"
        )
        for p in todos_pacientes:
          p_nombre = p["nombre"]
          p_id = p.get("id")
          
          ya_marcado_hoy = p_nombre in datos[fecha_str]

          cols_sel = st.columns([3, 2])
          seleccionado = cols_sel[0].checkbox(
              f"{p_nombre} ({p.get('turno','05:40')})",
              value=ya_marcado_hoy,
              key=f"chk_sel_{fecha_str}_{p_id}",
          )

          turno_actual_dia = (
              datos[fecha_str][p_nombre].get(
                  "turno_dia", p.get("turno", "05:40")
              )
              if ya_marcado_hoy
              else p.get("turno", "05:40")
          )
          idx_turno = (
              TURNOS_DISPONIBLES.index(turno_actual_dia)
              if turno_actual_dia in TURNOS_DISPONIBLES
              else 0
          )

          turno_elegido = cols_sel[1].selectbox(
              "Turno hoy",
              TURNOS_DISPONIBLES,
              index=idx_turno,
              key=f"sel_trn_{fecha_str}_{p_id}",
              label_visibility="collapsed",
          )

          if seleccionado:
            if p_nombre not in datos[fecha_str]:
              datos[fecha_str][p_nombre] = {
                  "ida": False,
                  "vuelta": False,
                  "listo": False,
                  "estado": "Ausente (No viajó)",
                  "valor": 0,
                  "turno_dia": turno_elegido,
              }
            else:
              datos[fecha_str][p_nombre]["turno_dia"] = turno_elegido
            guardar_registro_db(fecha_str, p_nombre, datos[fecha_str][p_nombre])
          else:
            if p_nombre in datos[fecha_str]:
              del datos[fecha_str][p_nombre]
              eliminar_registro_db(fecha_str, p_nombre)

      pacientes_hoy_dict = datos.get(fecha_str, {})

      def esta_finalizado(p_nombre):
        info = pacientes_hoy_dict.get(p_nombre, {})
        return (info.get("ida", False) and info.get("vuelta", False)) or info.get(
            "listo", False
        )

      pendientes_totales = sum(
          1 for nom in pacientes_hoy_dict if not esta_finalizado(nom)
      )
      st.metric(label="👥 Pacientes Pendientes Hoy", value=pendientes_totales)

      # --- RENDERIZAR TURNOS EXCLUSIVAMENTE CON LOS SELECCIONADOS DEL DÍA ---
      for turno_turno, titulo_turno in [
          ("05:40", "⏰ Turno 05:40 AM"),
          ("08:40", "⏰ Turno 08:40 AM"),
      ]:
        st.markdown(f"### {titulo_turno}")
        
        pacientes_turno = []
        for p in todos_pacientes:
          nombre_p = p["nombre"]
          if nombre_p in pacientes_hoy_dict:
            if pacientes_hoy_dict[nombre_p].get("turno_dia", "05:40") == turno_turno:
              pacientes_turno.append(p)

        if not pacientes_turno:
          st.write("*(Sin pacientes programados en este turno)*")
        else:
          pacientes_turno.sort(
              key=lambda p: (
                  1 if pacientes_hoy_dict[p["nombre"]].get("listo", False) else 0,
                  p["nombre"]
              )
          )

        for p in pacientes_turno:
          nombre = p["nombre"]
          p_id = p["id"]
          info_paciente = pacientes_hoy_dict[nombre]
          reg_ida = info_paciente.get("ida", False)
          reg_vuelta = info_paciente.get("vuelta", False)
          is_listo = info_paciente.get("listo", False)

          with st.container(border=True):
            estado_icono = (
                "🟢"
                if (reg_ida and reg_vuelta)
                else ("🟡" if (reg_ida or reg_vuelta or is_listo) else "⚪")
            )
            finalizado_txt = " *[Finalizado]*" if is_listo else ""

            with st.expander(
                f"{estado_icono} {nombre}{finalizado_txt}", expanded=False
            ):
              st.markdown(f"**Editando datos de {nombre}:**")
              edit_nombre = st.text_input(
                  "Nombre y Apellido", value=p["nombre"], key=f"edit_nom_{p_id}"
              )
              edit_grupo = st.selectbox(
                  "Grupo Asignado",
                  [
                      "Grupo 1 (Lunes, Miércoles, Viernes)",
                      "Grupo 2 (Martes, Jueves, Sábados)",
                  ],
                  index=0 if p["grupo"] == "MWF" else 1,
                  key=f"edit_grp_{p_id}",
              )
              edit_turno = st.selectbox(
                  "Turno Habitual",
                  TURNOS_DISPONIBLES,
                  index=(
                      TURNOS_DISPONIBLES.index(p["turno"])
                      if p["turno"] in TURNOS_DISPONIBLES
                      else 0
                  ),
                  key=f"
